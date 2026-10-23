#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
WOS Teacher vs Distilled Student Case Analysis
============================================================

目标：
1. 找出 Teacher 回答错误，但蒸馏后的 Student 回答正确的具体 case。
2. 保留原始正文、Gold、Teacher 预测、Student 预测。
3. 按 L1 / L2 错误类型分类。
4. 最终直接生成 Excel (.xlsx)。

当前默认比较：

Teacher:
    Llama-3.1-8B SFT Teacher

Student:
    Full KD
    Align = 1
    Hier  = 0.001

核心定义：
    Teacher 至少有一层错误
    AND
    Student L1 正确
    AND
    Student L2 正确

最终 Excel Sheet：
    1. Summary
    2. Teacher错_Student全对      <-- 最重要
    3. L2修正                    <-- Teacher L1对、L2错
    4. L1修正                    <-- Teacher L1错、L2对
    5. L1L2全修正                <-- Teacher L1/L2都错
    6. L1单层修正
    7. L2单层修正
    8. Teacher对_Student错
    9. 全部样本
"""

import os
import json
import argparse
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


# =============================================================================
# 1. 默认文件路径
# =============================================================================

# Teacher 测试生成结果
DEFAULT_TEACHER_FILE = (
    "/data1/lzp/code/hierarchical_distill/"
    "checkpoints_teacher_wos/test_results_best/"
    "teacher_test_generations.json"
)

# 最终蒸馏模型：
# Align = 1
# Hier  = 0.001
DEFAULT_STUDENT_FILE = (
    "/data1/lzp/code/hierarchical_distill/"
    "checkpoints_wos_full_kd_hier0001/test_results_best/"
    "offline_traditional_kd_test_generations.json"
)

# WOS TEST 原始数据
DEFAULT_TEST_FILE = (
    "/data1/lzp/code/data/wos/"
    "test_formatted.jsonl"
)

# 输出 Excel
DEFAULT_OUTPUT_FILE = (
    "/data1/lzp/code/hierarchical_distill/WOS/"
    "WOS_teacher_vs_full_hier0001_case_analysis.xlsx"
)


# =============================================================================
# 2. Excel 最大单元格字符数
# =============================================================================

EXCEL_CELL_LIMIT = 32767


# =============================================================================
# 3. 通用读文件
# =============================================================================

def load_records(path):
    """
    同时支持：
        .json
        .jsonl

    JSON 支持：
        [...]
        {"predictions": [...]}
        {"results": [...]}
        {"generations": [...]}
        {"data": [...]}
    """

    if not os.path.exists(path):
        raise FileNotFoundError(
            f"\n❌ 找不到文件：\n{path}\n"
        )

    # -------------------------------------------------------------------------
    # JSONL
    # -------------------------------------------------------------------------

    if path.lower().endswith(".jsonl"):

        rows = []

        with open(
            path,
            "r",
            encoding="utf-8",
        ) as f:

            for line in f:

                line = line.strip()

                if not line:
                    continue

                rows.append(
                    json.loads(line)
                )

        return rows

    # -------------------------------------------------------------------------
    # JSON
    # -------------------------------------------------------------------------

    with open(
        path,
        "r",
        encoding="utf-8",
    ) as f:

        obj = json.load(f)

    if isinstance(obj, list):
        return obj

    if isinstance(obj, dict):

        possible_keys = [
            "predictions",
            "results",
            "generations",
            "data",
        ]

        for key in possible_keys:

            if (
                key in obj
                and isinstance(obj[key], list)
            ):
                return obj[key]

    raise ValueError(
        f"\n❌ 无法识别 JSON 数据结构：\n{path}\n"
    )


# =============================================================================
# 4. 兼容多个字段名称
# =============================================================================

def get_field(
    row,
    possible_keys,
    required=True,
    default=None,
):

    for key in possible_keys:

        if key in row:
            return row[key]

    if required:

        raise KeyError(
            "\n❌ 找不到需要的字段\n"
            f"候选字段：{possible_keys}\n"
            f"当前已有字段：{list(row.keys())}\n"
        )

    return default


# =============================================================================
# 5. bool 转换
# =============================================================================

def to_bool(x):

    if isinstance(x, bool):
        return x

    if isinstance(x, (int, float)):
        return bool(x)

    if isinstance(x, str):

        x = x.strip().lower()

        if x in {
            "true",
            "1",
            "yes",
            "y",
        }:
            return True

        if x in {
            "false",
            "0",
            "no",
            "n",
            "",
        }:
            return False

    return bool(x)


# =============================================================================
# 6. Excel 安全文本
# =============================================================================

def excel_safe(x):
    """
    1. None -> ""
    2. 防止以 = 开头的正文被 Excel 当作公式
    3. 防止超过 Excel 单元格 32767 字符限制
    """

    if x is None:
        return ""

    text = str(x)

    if text.startswith("="):
        text = "'" + text

    if len(text) > EXCEL_CELL_LIMIT:

        text = (
            text[:EXCEL_CELL_LIMIT - 30]
            + "\n...[TRUNCATED FOR EXCEL]"
        )

    return text


# =============================================================================
# 7. 提取统一的预测格式
# =============================================================================

def parse_prediction(row):

    source_index = int(
        get_field(
            row,
            [
                "source_index",
                "index",
            ],
        )
    )

    true_l1 = excel_safe(
        get_field(
            row,
            [
                "true_l1",
                "gold_l1",
            ],
        )
    )

    true_l2 = excel_safe(
        get_field(
            row,
            [
                "true_l2",
                "gold_l2",
            ],
        )
    )

    pred_l1 = excel_safe(
        get_field(
            row,
            [
                "extracted_l1_exact",
                "pred_l1",
                "prediction_l1",
            ],
        )
    )

    pred_l2 = excel_safe(
        get_field(
            row,
            [
                "extracted_l2_exact",
                "pred_l2",
                "prediction_l2",
            ],
        )
    )

    correct_l1 = to_bool(
        get_field(
            row,
            [
                "correct_l1",
            ],
        )
    )

    correct_l2 = to_bool(
        get_field(
            row,
            [
                "correct_l2",
            ],
        )
    )

    raw_output = excel_safe(
        get_field(
            row,
            [
                "raw_generated_text_with_special_tokens",
                "raw_generated_text",
                "generated_text",
            ],
            required=False,
            default="",
        )
    )

    return {
        "source_index": source_index,

        "true_l1": true_l1,
        "true_l2": true_l2,

        "pred_l1": pred_l1,
        "pred_l2": pred_l2,

        "correct_l1": correct_l1,
        "correct_l2": correct_l2,

        "raw_output": raw_output,
    }


# =============================================================================
# 8. 判断 case 类型
# =============================================================================

def determine_case_type(
    t_l1,
    t_l2,
    s_l1,
    s_l2,
):

    teacher_both = t_l1 and t_l2
    student_both = s_l1 and s_l2

    # -------------------------------------------------------------------------
    # Teacher 错，Student 全对
    # -------------------------------------------------------------------------

    if (
        not teacher_both
        and student_both
    ):

        if (
            not t_l1
            and not t_l2
        ):

            return (
                "Teacher_L1L2都错_"
                "Student全部修正"
            )

        if (
            not t_l1
            and t_l2
        ):

            return (
                "Teacher_L1错_"
                "Student修正"
            )

        if (
            t_l1
            and not t_l2
        ):

            return (
                "Teacher_L2错_"
                "Student修正"
            )

    # -------------------------------------------------------------------------
    # Teacher 全对，Student 错
    # -------------------------------------------------------------------------

    if (
        teacher_both
        and not student_both
    ):

        if (
            not s_l1
            and not s_l2
        ):

            return (
                "Teacher正确_"
                "Student_L1L2都错"
            )

        if not s_l1:

            return (
                "Teacher正确_"
                "Student_L1错"
            )

        if not s_l2:

            return (
                "Teacher正确_"
                "Student_L2错"
            )

    # -------------------------------------------------------------------------
    # 都对
    # -------------------------------------------------------------------------

    if (
        teacher_both
        and student_both
    ):

        return (
            "Teacher正确_Student正确"
        )

    # -------------------------------------------------------------------------
    # 两者都有错误
    # -------------------------------------------------------------------------

    return (
        "Teacher和Student都存在错误"
    )


# =============================================================================
# 9. 建立所有 case
# =============================================================================

def build_all_cases(
    teacher_results,
    student_results,
    test_data,
):

    # -------------------------------------------------------------------------
    # Teacher index
    # -------------------------------------------------------------------------

    teacher_by_idx = {}

    for row in teacher_results:

        parsed = parse_prediction(row)

        idx = parsed["source_index"]

        if idx in teacher_by_idx:

            raise ValueError(
                f"Teacher source_index 重复：{idx}"
            )

        teacher_by_idx[idx] = parsed

    # -------------------------------------------------------------------------
    # Student index
    # -------------------------------------------------------------------------

    student_by_idx = {}

    for row in student_results:

        parsed = parse_prediction(row)

        idx = parsed["source_index"]

        if idx in student_by_idx:

            raise ValueError(
                f"Student source_index 重复：{idx}"
            )

        student_by_idx[idx] = parsed

    # -------------------------------------------------------------------------
    # 对齐
    # -------------------------------------------------------------------------

    teacher_indices = set(
        teacher_by_idx.keys()
    )

    student_indices = set(
        student_by_idx.keys()
    )

    if teacher_indices != student_indices:

        teacher_only = sorted(
            teacher_indices
            - student_indices
        )

        student_only = sorted(
            student_indices
            - teacher_indices
        )

        raise ValueError(
            "\n❌ Teacher / Student source_index 不一致\n"
            f"Teacher only 前20个：{teacher_only[:20]}\n"
            f"Student only 前20个：{student_only[:20]}\n"
        )

    print(
        "✅ Teacher / Student source_index 完全一致"
    )

    # -------------------------------------------------------------------------
    # TEST 数据建立 source_index 映射
    # -------------------------------------------------------------------------

    raw_by_idx = {}

    for i, row in enumerate(test_data):

        if "source_index" in row:

            idx = int(
                row["source_index"]
            )

        else:

            # 当前 WOS test_formatted.jsonl
            # source_index 就对应行号
            idx = i

        raw_by_idx[idx] = row

    # -------------------------------------------------------------------------
    # 构建
    # -------------------------------------------------------------------------

    all_cases = []

    for idx in sorted(teacher_indices):

        if idx not in raw_by_idx:

            raise ValueError(
                f"TEST 中找不到 source_index={idx}"
            )

        t = teacher_by_idx[idx]
        s = student_by_idx[idx]
        raw = raw_by_idx[idx]

        # ---------------------------------------------------------------------
        # Teacher / Student Gold 必须一致
        # ---------------------------------------------------------------------

        if (
            t["true_l1"]
            != s["true_l1"]
        ):

            raise ValueError(
                f"\nindex={idx}\n"
                "Teacher / Student Gold L1 不一致\n"
                f"Teacher: {t['true_l1']}\n"
                f"Student: {s['true_l1']}"
            )

        if (
            t["true_l2"]
            != s["true_l2"]
        ):

            raise ValueError(
                f"\nindex={idx}\n"
                "Teacher / Student Gold L2 不一致\n"
                f"Teacher: {t['true_l2']}\n"
                f"Student: {s['true_l2']}"
            )

        gold_l1 = t["true_l1"]
        gold_l2 = t["true_l2"]

        # ---------------------------------------------------------------------
        # TEST Gold 再检查一遍
        # ---------------------------------------------------------------------

        raw_l1 = excel_safe(
            get_field(
                raw,
                [
                    "l1_label",
                    "true_l1",
                    "gold_l1",
                ],
            )
        )

        raw_l2 = excel_safe(
            get_field(
                raw,
                [
                    "l2_label",
                    "true_l2",
                    "gold_l2",
                ],
            )
        )

        if raw_l1 != gold_l1:

            raise ValueError(
                f"\n❌ index={idx} L1 Gold 对不上\n"
                f"TEST       : {raw_l1}\n"
                f"Prediction : {gold_l1}\n"
            )

        if raw_l2 != gold_l2:

            raise ValueError(
                f"\n❌ index={idx} L2 Gold 对不上\n"
                f"TEST       : {raw_l2}\n"
                f"Prediction : {gold_l2}\n"
            )

        # ---------------------------------------------------------------------
        # 正文
        # ---------------------------------------------------------------------

        text = excel_safe(
            get_field(
                raw,
                [
                    "text",
                    "content",
                    "document",
                ],
            )
        )

        # ---------------------------------------------------------------------
        # Correct
        # ---------------------------------------------------------------------

        t_l1 = t["correct_l1"]
        t_l2 = t["correct_l2"]

        s_l1 = s["correct_l1"]
        s_l2 = s["correct_l2"]

        t_both = (
            t_l1
            and t_l2
        )

        s_both = (
            s_l1
            and s_l2
        )

        case_type = determine_case_type(
            t_l1=t_l1,
            t_l2=t_l2,
            s_l1=s_l1,
            s_l2=s_l2,
        )

        # ---------------------------------------------------------------------
        # Teacher 错误类型
        # ---------------------------------------------------------------------

        if (
            not t_l1
            and not t_l2
        ):

            teacher_error_type = (
                "L1+L2 都错"
            )

        elif not t_l1:

            teacher_error_type = (
                "只有 L1 错"
            )

        elif not t_l2:

            teacher_error_type = (
                "只有 L2 错"
            )

        else:

            teacher_error_type = (
                "无错误"
            )

        # ---------------------------------------------------------------------
        # Student 错误类型
        # ---------------------------------------------------------------------

        if (
            not s_l1
            and not s_l2
        ):

            student_error_type = (
                "L1+L2 都错"
            )

        elif not s_l1:

            student_error_type = (
                "只有 L1 错"
            )

        elif not s_l2:

            student_error_type = (
                "只有 L2 错"
            )

        else:

            student_error_type = (
                "无错误"
            )

        all_cases.append({
            "source_index": idx,

            "case_type":
                case_type,

            "teacher_error_type":
                teacher_error_type,

            "student_error_type":
                student_error_type,

            # Gold
            "gold_l1":
                gold_l1,

            "gold_l2":
                gold_l2,

            # Teacher
            "teacher_pred_l1":
                t["pred_l1"],

            "teacher_pred_l2":
                t["pred_l2"],

            "teacher_l1_correct":
                t_l1,

            "teacher_l2_correct":
                t_l2,

            "teacher_both_correct":
                t_both,

            # Student
            "student_pred_l1":
                s["pred_l1"],

            "student_pred_l2":
                s["pred_l2"],

            "student_l1_correct":
                s_l1,

            "student_l2_correct":
                s_l2,

            "student_both_correct":
                s_both,

            # 变化
            "l1_change":
                (
                    f"{t['pred_l1']}"
                    f"  →  "
                    f"{s['pred_l1']}"
                ),

            "l2_change":
                (
                    f"{t['pred_l2']}"
                    f"  →  "
                    f"{s['pred_l2']}"
                ),

            # 正文
            "text":
                text,

            # 原始生成
            "teacher_raw_output":
                t["raw_output"],

            "student_raw_output":
                s["raw_output"],
        })

    return all_cases


# =============================================================================
# 10. Excel 列
# =============================================================================

HEADERS = [
    "Case No.",
    "Test Index",
    "Case Type",

    "Teacher Error Type",
    "Student Error Type",

    "Gold L1",
    "Gold L2",

    "Teacher L1",
    "Teacher L2",

    "Teacher L1 Correct",
    "Teacher L2 Correct",

    "Student L1",
    "Student L2",

    "Student L1 Correct",
    "Student L2 Correct",

    "L1: Teacher → Student",
    "L2: Teacher → Student",

    "Original Text",

    "Teacher Raw Output",
    "Student Raw Output",
]


def case_to_row(
    case_no,
    case,
):

    return [
        case_no,
        case["source_index"],
        case["case_type"],

        case["teacher_error_type"],
        case["student_error_type"],

        case["gold_l1"],
        case["gold_l2"],

        case["teacher_pred_l1"],
        case["teacher_pred_l2"],

        case["teacher_l1_correct"],
        case["teacher_l2_correct"],

        case["student_pred_l1"],
        case["student_pred_l2"],

        case["student_l1_correct"],
        case["student_l2_correct"],

        case["l1_change"],
        case["l2_change"],

        case["text"],

        case["teacher_raw_output"],
        case["student_raw_output"],
    ]


# =============================================================================
# 11. Excel 样式
# =============================================================================

HEADER_FILL = PatternFill(
    "solid",
    fgColor="1F4E78",
)

HEADER_FONT = Font(
    color="FFFFFF",
    bold=True,
)

GOLD_FILL = PatternFill(
    "solid",
    fgColor="FFF2CC",
)

TEACHER_WRONG_FILL = PatternFill(
    "solid",
    fgColor="F4CCCC",
)

TEACHER_WRONG_FONT = Font(
    color="9C0006",
    bold=True,
)

STUDENT_CORRECT_FILL = PatternFill(
    "solid",
    fgColor="C6EFCE",
)

STUDENT_CORRECT_FONT = Font(
    color="006100",
    bold=True,
)

STUDENT_WRONG_FILL = PatternFill(
    "solid",
    fgColor="F4CCCC",
)

THIN_GRAY = Side(
    style="thin",
    color="D9E1F2",
)

CELL_BORDER = Border(
    left=THIN_GRAY,
    right=THIN_GRAY,
    top=THIN_GRAY,
    bottom=THIN_GRAY,
)


# =============================================================================
# 12. 添加 case Sheet
# =============================================================================

def add_case_sheet(
    wb,
    sheet_name,
    cases,
    table_name,
):

    ws = wb.create_sheet(
        title=sheet_name
    )

    # -------------------------------------------------------------------------
    # Header
    # -------------------------------------------------------------------------

    for col_idx, header in enumerate(
        HEADERS,
        start=1,
    ):

        cell = ws.cell(
            row=1,
            column=col_idx,
            value=header,
        )

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        cell.border = CELL_BORDER

    ws.row_dimensions[1].height = 30

    # -------------------------------------------------------------------------
    # 数据
    # -------------------------------------------------------------------------

    for row_idx, case in enumerate(
        cases,
        start=2,
    ):

        values = case_to_row(
            row_idx - 1,
            case,
        )

        for col_idx, value in enumerate(
            values,
            start=1,
        ):

            cell = ws.cell(
                row=row_idx,
                column=col_idx,
                value=value,
            )

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            cell.border = CELL_BORDER

        # ---------------------------------------------------------------------
        # Gold
        #
        # F = Gold L1
        # G = Gold L2
        # ---------------------------------------------------------------------

        ws.cell(
            row_idx,
            6,
        ).fill = GOLD_FILL

        ws.cell(
            row_idx,
            7,
        ).fill = GOLD_FILL

        # ---------------------------------------------------------------------
        # Teacher
        #
        # H Teacher L1
        # I Teacher L2
        # J Teacher L1 correct
        # K Teacher L2 correct
        # ---------------------------------------------------------------------

        if not case["teacher_l1_correct"]:

            ws.cell(
                row_idx,
                8,
            ).fill = TEACHER_WRONG_FILL

            ws.cell(
                row_idx,
                8,
            ).font = TEACHER_WRONG_FONT

            ws.cell(
                row_idx,
                10,
            ).fill = TEACHER_WRONG_FILL

        if not case["teacher_l2_correct"]:

            ws.cell(
                row_idx,
                9,
            ).fill = TEACHER_WRONG_FILL

            ws.cell(
                row_idx,
                9,
            ).font = TEACHER_WRONG_FONT

            ws.cell(
                row_idx,
                11,
            ).fill = TEACHER_WRONG_FILL

        # ---------------------------------------------------------------------
        # Student
        #
        # L Student L1
        # M Student L2
        # N Student L1 correct
        # O Student L2 correct
        # ---------------------------------------------------------------------

        if case["student_l1_correct"]:

            ws.cell(
                row_idx,
                12,
            ).fill = STUDENT_CORRECT_FILL

            ws.cell(
                row_idx,
                12,
            ).font = STUDENT_CORRECT_FONT

            ws.cell(
                row_idx,
                14,
            ).fill = STUDENT_CORRECT_FILL

        else:

            ws.cell(
                row_idx,
                12,
            ).fill = STUDENT_WRONG_FILL

            ws.cell(
                row_idx,
                14,
            ).fill = STUDENT_WRONG_FILL

        if case["student_l2_correct"]:

            ws.cell(
                row_idx,
                13,
            ).fill = STUDENT_CORRECT_FILL

            ws.cell(
                row_idx,
                13,
            ).font = STUDENT_CORRECT_FONT

            ws.cell(
                row_idx,
                15,
            ).fill = STUDENT_CORRECT_FILL

        else:

            ws.cell(
                row_idx,
                13,
            ).fill = STUDENT_WRONG_FILL

            ws.cell(
                row_idx,
                15,
            ).fill = STUDENT_WRONG_FILL

        # 不让正文把每一行直接撑得非常高
        ws.row_dimensions[row_idx].height = 55

    # -------------------------------------------------------------------------
    # 冻结
    #
    # C2：
    #   冻结第一行
    #   冻结前两列 Case No / Test Index
    # -------------------------------------------------------------------------

    ws.freeze_panes = "C2"

    # -------------------------------------------------------------------------
    # Filter
    # -------------------------------------------------------------------------

    if len(cases) > 0:

        last_row = len(cases) + 1
        last_col = len(HEADERS)

        ws.auto_filter.ref = (
            f"A1:"
            f"{get_column_letter(last_col)}"
            f"{last_row}"
        )

    # -------------------------------------------------------------------------
    # 列宽
    # -------------------------------------------------------------------------

    widths = {
        "A": 10,     # Case No
        "B": 12,     # Test Index
        "C": 34,     # Case Type

        "D": 18,     # Teacher Error
        "E": 18,     # Student Error

        "F": 22,     # Gold L1
        "G": 30,     # Gold L2

        "H": 22,     # Teacher L1
        "I": 30,     # Teacher L2

        "J": 18,
        "K": 18,

        "L": 22,     # Student L1
        "M": 30,     # Student L2

        "N": 18,
        "O": 18,

        "P": 40,
        "Q": 40,

        "R": 65,     # Original text

        "S": 42,
        "T": 42,
    }

    for col, width in widths.items():

        ws.column_dimensions[
            col
        ].width = width

    # -------------------------------------------------------------------------
    # 对齐
    # -------------------------------------------------------------------------

    for row in ws.iter_rows(
        min_row=2,
        max_row=len(cases) + 1,
        min_col=1,
        max_col=17,
    ):

        for cell in row:

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

    # Original Text 和 Raw Output 左对齐
    for row_idx in range(
        2,
        len(cases) + 2,
    ):

        for col_idx in [
            18,
            19,
            20,
        ]:

            ws.cell(
                row_idx,
                col_idx,
            ).alignment = Alignment(
                horizontal="left",
                vertical="top",
                wrap_text=True,
            )

    return ws


# =============================================================================
# 13. Summary Sheet
# =============================================================================

def create_summary_sheet(
    wb,
    all_cases,
    recovery,
    recovery_l2,
    recovery_l1,
    recovery_both,
    l1_recovery_any,
    l2_recovery_any,
    failure,
):

    ws = wb.active

    ws.title = "Summary"

    # -------------------------------------------------------------------------
    # 总体统计
    # -------------------------------------------------------------------------

    total = len(all_cases)

    teacher_all_correct = sum(
        1
        for c in all_cases
        if c["teacher_both_correct"]
    )

    student_all_correct = sum(
        1
        for c in all_cases
        if c["student_both_correct"]
    )

    teacher_l1_correct_student_wrong = sum(
        1
        for c in all_cases
        if (
            c["teacher_l1_correct"]
            and not c["student_l1_correct"]
        )
    )

    teacher_l2_correct_student_wrong = sum(
        1
        for c in all_cases
        if (
            c["teacher_l2_correct"]
            and not c["student_l2_correct"]
        )
    )

    summary = [
        [
            "WOS Teacher vs Distilled Student Case Analysis",
            "",
        ],

        [
            "统计项",
            "数量",
        ],

        [
            "TEST 总样本数",
            total,
        ],

        [
            "Teacher L1/L2 全对",
            teacher_all_correct,
        ],

        [
            "Student L1/L2 全对",
            student_all_correct,
        ],

        [
            "Teacher至少一层错 → Student两层全对",
            len(recovery),
        ],

        [
            "Teacher只有L2错 → Student两层全对",
            len(recovery_l2),
        ],

        [
            "Teacher只有L1错 → Student两层全对",
            len(recovery_l1),
        ],

        [
            "Teacher L1/L2都错 → Student两层全对",
            len(recovery_both),
        ],

        [
            "Teacher两层全对 → Student至少一层错",
            len(failure),
        ],

        [
            "两层全部正确净收益",
            (
                len(recovery)
                - len(failure)
            ),
        ],

        [
            "Teacher L1错 → Student L1对",
            len(l1_recovery_any),
        ],

        [
            "Teacher L1对 → Student L1错",
            teacher_l1_correct_student_wrong,
        ],

        [
            "L1净修正",
            (
                len(l1_recovery_any)
                - teacher_l1_correct_student_wrong
            ),
        ],

        [
            "Teacher L2错 → Student L2对",
            len(l2_recovery_any),
        ],

        [
            "Teacher L2对 → Student L2错",
            teacher_l2_correct_student_wrong,
        ],

        [
            "L2净修正",
            (
                len(l2_recovery_any)
                - teacher_l2_correct_student_wrong
            ),
        ],
    ]

    for r, row in enumerate(
        summary,
        start=1,
    ):

        for c, value in enumerate(
            row,
            start=1,
        ):

            ws.cell(
                r,
                c,
                value,
            )

    # -------------------------------------------------------------------------
    # 标题
    # -------------------------------------------------------------------------

    ws.merge_cells(
        "A1:B1"
    )

    ws["A1"].fill = HEADER_FILL
    ws["A1"].font = Font(
        color="FFFFFF",
        bold=True,
        size=15,
    )

    ws["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    ws.row_dimensions[1].height = 30

    # Header
    for cell in ws[2]:

        cell.fill = PatternFill(
            "solid",
            fgColor="5B9BD5",
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.alignment = Alignment(
            horizontal="center",
        )

    # 恢复数量绿色
    for row_idx in range(
        6,
        10,
    ):

        for col_idx in range(
            1,
            3,
        ):

            ws.cell(
                row_idx,
                col_idx,
            ).fill = PatternFill(
                "solid",
                fgColor="E2F0D9",
            )

    # 失败红色
    for col_idx in range(
        1,
        3,
    ):

        ws.cell(
            10,
            col_idx,
        ).fill = PatternFill(
            "solid",
            fgColor="F4CCCC",
        )

    # 净收益黄色
    for col_idx in range(
        1,
        3,
    ):

        ws.cell(
            11,
            col_idx,
        ).fill = PatternFill(
            "solid",
            fgColor="FFF2CC",
        )

    # -------------------------------------------------------------------------
    # 最常被救回的 Gold L1 / L2
    # -------------------------------------------------------------------------

    recovered_l1_counter = Counter(
        c["gold_l1"]
        for c in recovery
    )

    recovered_l2_counter = Counter(
        c["gold_l2"]
        for c in recovery
    )

    ws["D2"] = (
        "Student 最常救回的 Gold L1"
    )

    ws["E2"] = "数量"

    ws["G2"] = (
        "Student 最常救回的 Gold L2"
    )

    ws["H2"] = "数量"

    for cell in [
        ws["D2"],
        ws["E2"],
        ws["G2"],
        ws["H2"],
    ]:

        cell.fill = PatternFill(
            "solid",
            fgColor="70AD47",
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.alignment = Alignment(
            horizontal="center",
        )

    for i, (
        label,
        count,
    ) in enumerate(
        recovered_l1_counter.most_common(20),
        start=3,
    ):

        ws.cell(
            i,
            4,
            label,
        )

        ws.cell(
            i,
            5,
            count,
        )

    for i, (
        label,
        count,
    ) in enumerate(
        recovered_l2_counter.most_common(30),
        start=3,
    ):

        ws.cell(
            i,
            7,
            label,
        )

        ws.cell(
            i,
            8,
            count,
        )

    # -------------------------------------------------------------------------
    # Sheet 说明
    # -------------------------------------------------------------------------

    start_row = 36

    notes = [
        [
            "Sheet",
            "说明",
        ],

        [
            "Teacher错_Student全对",
            (
                "最核心：Teacher至少有一层错误，"
                "Student L1和L2都正确"
            ),
        ],

        [
            "L2修正",
            (
                "Teacher L1正确、L2错误，"
                "Student L1/L2都正确"
            ),
        ],

        [
            "L1修正",
            (
                "Teacher L1错误、L2正确，"
                "Student L1/L2都正确"
            ),
        ],

        [
            "L1L2全修正",
            (
                "Teacher L1和L2都错误，"
                "Student全部修正"
            ),
        ],

        [
            "L1单层修正",
            (
                "只看L1：Teacher L1错，"
                "Student L1对"
            ),
        ],

        [
            "L2单层修正",
            (
                "只看L2：Teacher L2错，"
                "Student L2对"
            ),
        ],

        [
            "Teacher对_Student错",
            (
                "反向失败案例：Teacher两层正确，"
                "Student至少一层错误"
            ),
        ],

        [
            "全部样本",
            "全部测试样本逐条对照",
        ],
    ]

    for r_offset, row in enumerate(
        notes,
    ):

        for c_offset, value in enumerate(
            row,
        ):

            ws.cell(
                start_row + r_offset,
                1 + c_offset,
                value,
            )

    ws.cell(
        start_row,
        1,
    ).fill = HEADER_FILL

    ws.cell(
        start_row,
        2,
    ).fill = HEADER_FILL

    ws.cell(
        start_row,
        1,
    ).font = HEADER_FONT

    ws.cell(
        start_row,
        2,
    ).font = HEADER_FONT

    # -------------------------------------------------------------------------
    # Width
    # -------------------------------------------------------------------------

    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 18

    ws.column_dimensions["D"].width = 34
    ws.column_dimensions["E"].width = 12

    ws.column_dimensions["G"].width = 42
    ws.column_dimensions["H"].width = 12

    ws.freeze_panes = "A3"

    for row in ws.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

    return ws


# =============================================================================
# 14. 主程序
# =============================================================================

def main():

    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--teacher_file",
        type=str,
        default=DEFAULT_TEACHER_FILE,
    )

    parser.add_argument(
        "--student_file",
        type=str,
        default=DEFAULT_STUDENT_FILE,
    )

    parser.add_argument(
        "--test_file",
        type=str,
        default=DEFAULT_TEST_FILE,
    )

    parser.add_argument(
        "--output",
        type=str,
        default=DEFAULT_OUTPUT_FILE,
    )

    args = parser.parse_args()

    print("=" * 100)
    print(
        "WOS Teacher vs Distilled Student Case Analysis"
    )
    print("=" * 100)

    print("\nTeacher:")
    print(args.teacher_file)

    print("\nStudent:")
    print(args.student_file)

    print("\nTEST:")
    print(args.test_file)

    print("\nOutput:")
    print(args.output)

    # =========================================================================
    # 读取
    # =========================================================================

    teacher_results = load_records(
        args.teacher_file
    )

    student_results = load_records(
        args.student_file
    )

    test_data = load_records(
        args.test_file
    )

    print("\n" + "=" * 100)
    print("读取数据")
    print("=" * 100)

    print(
        "Teacher samples:",
        len(teacher_results),
    )

    print(
        "Student samples:",
        len(student_results),
    )

    print(
        "TEST samples:",
        len(test_data),
    )

    if (
        len(teacher_results)
        != len(student_results)
    ):

        raise ValueError(
            "\n❌ Teacher / Student 样本数量不同"
        )

    # =========================================================================
    # 建所有 case
    # =========================================================================

    all_cases = build_all_cases(
        teacher_results,
        student_results,
        test_data,
    )

    # =========================================================================
    # 核心：
    #
    # Teacher 至少一层错
    # Student 两层都对
    # =========================================================================

    recovery = [
        c
        for c in all_cases
        if (
            not c["teacher_both_correct"]
            and c["student_both_correct"]
        )
    ]

    # -------------------------------------------------------------------------
    # Teacher L1对 L2错
    # -------------------------------------------------------------------------

    recovery_l2 = [
        c
        for c in recovery
        if (
            c["teacher_l1_correct"]
            and not c["teacher_l2_correct"]
        )
    ]

    # -------------------------------------------------------------------------
    # Teacher L1错 L2对
    # -------------------------------------------------------------------------

    recovery_l1 = [
        c
        for c in recovery
        if (
            not c["teacher_l1_correct"]
            and c["teacher_l2_correct"]
        )
    ]

    # -------------------------------------------------------------------------
    # Teacher L1/L2都错
    # -------------------------------------------------------------------------

    recovery_both = [
        c
        for c in recovery
        if (
            not c["teacher_l1_correct"]
            and not c["teacher_l2_correct"]
        )
    ]

    # -------------------------------------------------------------------------
    # 只看 L1
    # -------------------------------------------------------------------------

    l1_recovery_any = [
        c
        for c in all_cases
        if (
            not c["teacher_l1_correct"]
            and c["student_l1_correct"]
        )
    ]

    # -------------------------------------------------------------------------
    # 只看 L2
    # -------------------------------------------------------------------------

    l2_recovery_any = [
        c
        for c in all_cases
        if (
            not c["teacher_l2_correct"]
            and c["student_l2_correct"]
        )
    ]

    # -------------------------------------------------------------------------
    # Teacher全对 Student错
    # -------------------------------------------------------------------------

    failure = [
        c
        for c in all_cases
        if (
            c["teacher_both_correct"]
            and not c["student_both_correct"]
        )
    ]

    # =========================================================================
    # 打印数量
    # =========================================================================

    print("\n" + "=" * 100)
    print("核心统计")
    print("=" * 100)

    print(
        "\nTeacher 至少一层错 → Student 两层全对:",
        len(recovery),
    )

    print(
        "    Teacher 只有 L2 错:",
        len(recovery_l2),
    )

    print(
        "    Teacher 只有 L1 错:",
        len(recovery_l1),
    )

    print(
        "    Teacher L1/L2 都错:",
        len(recovery_both),
    )

    print(
        "\nTeacher 两层全对 → Student 至少一层错:",
        len(failure),
    )

    print(
        "\n两层全部正确净收益:",
        (
            len(recovery)
            - len(failure)
        ),
    )

    print(
        "\nTeacher L1错 → Student L1对:",
        len(l1_recovery_any),
    )

    print(
        "Teacher L2错 → Student L2对:",
        len(l2_recovery_any),
    )

    # =========================================================================
    # 创建 Excel
    # =========================================================================

    print("\n" + "=" * 100)
    print("开始生成 Excel")
    print("=" * 100)

    wb = Workbook()

    # Summary
    create_summary_sheet(
        wb=wb,

        all_cases=all_cases,

        recovery=recovery,

        recovery_l2=recovery_l2,

        recovery_l1=recovery_l1,

        recovery_both=recovery_both,

        l1_recovery_any=l1_recovery_any,

        l2_recovery_any=l2_recovery_any,

        failure=failure,
    )

    # -------------------------------------------------------------------------
    # 1. 最重要
    # -------------------------------------------------------------------------

    add_case_sheet(
        wb,
        "Teacher错_Student全对",
        recovery,
        "RecoveryAll",
    )

    # -------------------------------------------------------------------------
    # 2. L2 修正
    # -------------------------------------------------------------------------

    add_case_sheet(
        wb,
        "L2修正",
        recovery_l2,
        "RecoveryL2",
    )

    # -------------------------------------------------------------------------
    # 3. L1 修正
    # -------------------------------------------------------------------------

    add_case_sheet(
        wb,
        "L1修正",
        recovery_l1,
        "RecoveryL1",
    )

    # -------------------------------------------------------------------------
    # 4. 两层修正
    # -------------------------------------------------------------------------

    add_case_sheet(
        wb,
        "L1L2全修正",
        recovery_both,
        "RecoveryBoth",
    )

    # -------------------------------------------------------------------------
    # 5. 单独 L1
    # -------------------------------------------------------------------------

    add_case_sheet(
        wb,
        "L1单层修正",
        l1_recovery_any,
        "L1RecoveryAny",
    )

    # -------------------------------------------------------------------------
    # 6. 单独 L2
    # -------------------------------------------------------------------------

    add_case_sheet(
        wb,
        "L2单层修正",
        l2_recovery_any,
        "L2RecoveryAny",
    )

    # -------------------------------------------------------------------------
    # 7. 反向失败
    # -------------------------------------------------------------------------

    add_case_sheet(
        wb,
        "Teacher对_Student错",
        failure,
        "StudentFailure",
    )

    # -------------------------------------------------------------------------
    # 8. 所有样本
    # -------------------------------------------------------------------------

    add_case_sheet(
        wb,
        "全部样本",
        all_cases,
        "AllCases",
    )

    # =========================================================================
    # 保存
    # =========================================================================

    output_dir = os.path.dirname(
        args.output
    )

    if output_dir:

        os.makedirs(
            output_dir,
            exist_ok=True,
        )

    wb.save(
        args.output
    )

    print("\n" + "=" * 100)
    print("✅ 分析完成")
    print("=" * 100)

    print(
        "\nExcel 文件："
    )

    print(
        args.output
    )

    print(
        "\n最重要的 Sheet："
    )

    print(
        "Teacher错_Student全对"
    )

    print(
        "\n其中："
    )

    print(
        "L2修正 = Teacher L1正确、L2错误，Student全部正确"
    )

    print(
        "L1L2全修正 = Teacher两层都错，Student两层全部正确"
    )

    print("=" * 100)


if __name__ == "__main__":
    main()